import streamlit as st
import pandas as pd
import os
import openpyxl
import time
from datetime import datetime

# Configuração inicial da página
st.set_page_config(page_title="Gestão de Escala - Conviva", layout="wide")

# --- SISTEMA DE LOGIN MULTI-REGIÃO ---
def verificar_login():
    # Dicionário de regiões e senhas. Adicione ou altere conforme necessário.
    credenciais = {
        "APIAI": "apiai2026",
        "SANTOS": "santos2026",
        "BAURU": "bauru2026"
    }
    
    if "autenticado" not in st.session_state:
        st.session_state["autenticado"] = False
        st.session_state["regiao"] = None

    if not st.session_state["autenticado"]:
        st.title("🔒 Acesso Restrito - Conviva")
        st.write("Selecione a sua diretoria e digite a senha de acesso.")
        
        regiao_selecionada = st.selectbox("Região:", list(credenciais.keys()))
        senha_inserida = st.text_input("Senha:", type="password")
        
        if st.button("Entrar"):
            if senha_inserida == credenciais[regiao_selecionada]:
                st.session_state["autenticado"] = True
                st.session_state["regiao"] = regiao_selecionada
                st.rerun()
            else:
                st.error("Senha incorreta para a região selecionada. Tente novamente.")
        return False
    return True

# Apenas carrega o resto do sistema se o login for válido
if verificar_login():
    
    # Exibe no topo da tela qual região está logada
    st.sidebar.success(f"Logado na região: {st.session_state['regiao']}")
    
    # Nomes dos arquivos (Aqui você pode usar o st.session_state['regiao'] no futuro para carregar arquivos diferentes)
    ARQUIVO_BASE = "Mapeamento_GRUPO_06_APIAI_OUTUBRO ATUALIZADO.xlsx"
    ARQUIVO_EVENTOS = "historico_ocorrencias.csv"

    @st.cache_data
    def carregar_base():
        try:
            df = pd.read_excel(ARQUIVO_BASE, sheet_name="APIAI", header=6)
            df = df.dropna(subset=["Cuidadores"])
            return df
        except Exception as e:
            st.error(f"Erro ao carregar a planilha base: {e}")
            return pd.DataFrame()

    def carregar_eventos():
        if os.path.exists(ARQUIVO_EVENTOS):
            return pd.read_csv(ARQUIVO_EVENTOS)
        else:
            return pd.DataFrame(columns=["Data_Registro", "Data_Evento", "Escola", "Cuidador", "Tipo_Evento", "Observacao"])

    def salvar_evento(data_evento, escola, cuidador, tipo, obs):
        df_eventos = carregar_eventos()
        novo_evento = pd.DataFrame([{
            "Data_Registro": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "Data_Evento": data_evento.strftime("%d/%m/%Y"),
            "Escola": escola,
            "Cuidador": cuidador,
            "Tipo_Evento": tipo,
            "Observacao": obs
        }])
        df_eventos = pd.concat([df_eventos, novo_evento], ignore_index=True)
        df_eventos.to_csv(ARQUIVO_EVENTOS, index=False)
        return True

    def substituir_colaborador(escola, cuidador_antigo, novo_nome, novo_cpf):
        try:
            wb = openpyxl.load_workbook(ARQUIVO_BASE)
            ws = wb["APIAI"]
            
            for row_idx in range(8, ws.max_row + 1):
                cell_escola = ws.cell(row=row_idx, column=3).value
                cell_cuidador = ws.cell(row=row_idx, column=9).value
                
                if cell_escola == escola and cell_cuidador == cuidador_antigo:
                    ws.cell(row=row_idx, column=9).value = novo_nome.upper()
                    ws.cell(row=row_idx, column=21).value = novo_cpf
                    
                    obs_antiga = ws.cell(row=row_idx, column=22).value or ""
                    nova_obs = f"[{datetime.now().strftime('%d/%m')}] Substituiu {cuidador_antigo}. {obs_antiga}"
                    ws.cell(row=row_idx, column=22).value = nova_obs
                    
                    wb.save(ARQUIVO_BASE)
                    st.cache_data.clear()
                    return True
            return False
        except Exception as e:
            st.error(f"Erro ao editar o Excel: {e}")
            return False

    # --- INTERFACE PRINCIPAL ---
    col1, col2 = st.columns([8, 1])
    with col1:
        st.title("Gestão de Escala e Ocorrências")
    with col2:
        if st.button("Sair"):
            st.session_state["autenticado"] = False
            st.session_state["regiao"] = None
            st.rerun()

    menu = st.sidebar.selectbox("Navegação", ["Lançamento de Ocorrências", "Movimentação de Pessoal", "Relatório de Fechamento"])
    df_base = carregar_base()

    if not df_base.empty:
        if menu == "Lançamento de Ocorrências":
            st.header("Registro Diário de Ocorrências")
            
            escolas = df_base["Unidade Escolar"].dropna().unique().tolist()
            escola_selecionada = st.selectbox("Selecione a Unidade Escolar:", [""] + escolas)
            
            if escola_selecionada:
                cuidadores_escola = df_base[df_base["Unidade Escolar"] == escola_selecionada]["Cuidadores"].tolist()
                cuidador_selecionado = st.selectbox("Selecione o Colaborador:", [""] + cuidadores_escola)
                
                if cuidador_selecionado:
                    with st.form("form_evento", clear_on_submit=True):
                        col1, col2 = st.columns(2)
                        with col1:
                            data_evento = st.date_input("Data da Ocorrência:")
                        with col2:
                            tipo_evento = st.selectbox("Tipo de Evento:", ["Falta", "Atestado", "Diária", "Dobra"])
                        
                        observacao = st.text_area("Observação / Justificativa:")
                        submit = st.form_submit_button("Salvar Ocorrência")
                        
                        if submit:
                            salvar_evento(data_evento, escola_selecionada, cuidador_selecionado, tipo_evento, observacao)
                            st.success("Ocorrência salva com sucesso!")

        elif menu == "Movimentação de Pessoal":
            st.header("Desligamento e Nova Contratação")
            
            escolas = df_base["Unidade Escolar"].dropna().unique().tolist()
            escola_selecionada = st.selectbox("Selecione a Escola do Posto:", [""] + escolas)
            
            if escola_selecionada:
                cuidadores_escola = df_base[df_base["Unidade Escolar"] == escola_selecionada]["Cuidadores"].tolist()
                cuidador_antigo = st.selectbox("Colaborador a ser substituído (ou VAGO):", [""] + cuidadores_escola)
                
                if cuidador_antigo:
                    with st.form("form_substituicao"):
                        col1, col2 = st.columns(2)
                        with col1:
                            novo_nome = st.text_input("Nome Completo do Novo Colaborador:")
                        with col2:
                            novo_cpf = st.text_input("CPF do Novo Colaborador (apenas números):")
                        
                        submit_troca = st.form_submit_button("Confirmar Substituição")
                        
                        if submit_troca:
                            if novo_nome and novo_cpf:
                                sucesso = substituir_colaborador(escola_selecionada, cuidador_antigo, novo_nome, novo_cpf)
                                if sucesso:
                                    st.success(f"Substituição concluída! {novo_nome} agora ocupa o posto.")
                                    time.sleep(2)
                                    st.rerun()
                            else:
                                st.error("Preencha o Nome e o CPF.")

        elif menu == "Relatório de Fechamento":
            st.header("Fechamento do Mês")
            df_eventos = carregar_eventos()
            
            if not df_eventos.empty:
                st.subheader("Histórico Detalhado (Dia a Dia)")
                st.dataframe(df_eventos, use_container_width=True)
                
                csv_detalhado = df_eventos.to_csv(index=False).encode('utf-8')
                st.download_button("Baixar Histórico Detalhado (CSV)", data=csv_detalhado, file_name="Historico_Detalhado.csv", mime="text/csv")
                st.markdown("---")
                
                st.subheader("Resumo Somado por Colaborador")
                resumo_eventos = pd.crosstab(df_eventos["Cuidador"], df_eventos["Tipo_Evento"]).reset_index()
                df_fixo = df_base[["Cuidadores", "Unidade Escolar", "Carga\nHorária"]].rename(columns={"Cuidadores": "Cuidador", "Carga\nHorária": "Carga Horária"})
                relatorio_final = pd.merge(df_fixo, resumo_eventos, on="Cuidador", how="left").fillna(0)
                
                for col in ["Falta", "Atestado", "Diária", "Dobra"]:
                    if col in relatorio_final.columns:
                        relatorio_final[col] = relatorio_final[col].astype(int)
                
                st.dataframe(relatorio_final, use_container_width=True)
                csv_resumo = relatorio_final.to_csv(index=False).encode('utf-8')
                st.download_button("Baixar Resumo de Fechamento (CSV)", data=csv_resumo, file_name="Resumo_Pagamento.csv", mime="text/csv")
            else:
                st.info("Nenhuma ocorrência registrada até o momento.")
    else:
        st.warning("Não foi possível carregar os dados. Verifique a planilha.")